"""
从翻译对照表Excel生成语言包JSON
用法: python build_locale_from_excel.py <excel路径> <语言代码>
语言代码: en / ja / ko
"""
import json
import sys
import re
from openpyxl import load_workbook

def build_nested(flat_dict):
    """将展平的 key->value 字典重建为嵌套JSON结构"""
    result = {}
    for key, value in flat_dict.items():
        parts = []
        # 解析 a.b.c[0].d 这样的key
        tokens = re.split(r'\.|\[', key)
        for t in tokens:
            if t.endswith(']'):
                parts.append(('list', int(t[:-1])))
            else:
                parts.append(('dict', t))

        current = result
        for i, (ptype, pval) in enumerate(parts):
            is_last = (i == len(parts) - 1)
            if ptype == 'dict':
                if is_last:
                    current[pval] = value
                else:
                    if pval not in current:
                        # 预判下一级是list还是dict
                        next_type = parts[i+1][0]
                        current[pval] = [] if next_type == 'list' else {}
                    current = current[pval]
            elif ptype == 'list':
                if is_last:
                    # 确保列表足够长
                    while len(current) <= pval:
                        current.append(None)
                    current[pval] = value
                else:
                    while len(current) <= pval:
                        next_type = parts[i+1][0]
                        current.append([] if next_type == 'list' else {})
                    current = current[pval]
    return result

def main():
    if len(sys.argv) < 3:
        print('用法: python build_locale_from_excel.py <excel路径> <语言代码>')
        print('语言代码: en / ja / ko')
        sys.exit(1)

    excel_path = sys.argv[1]
    lang = sys.argv[2].lower()

    lang_col = {'en': 5, 'ja': 6, 'ko': 7}
    if lang not in lang_col:
        print(f'不支持的语言: {lang}，可选: en, ja, ko')
        sys.exit(1)

    col_idx = lang_col[lang]

    wb = load_workbook(excel_path)
    ws = wb['📝 翻译总表']

    flat = {}
    skipped = 0
    translated = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        key = row[2]  # C列: Key
        value = row[col_idx - 1]  # 翻译列

        if not key or not isinstance(key, str):
            continue
        if key.startswith('▸'):  # 分组标题行
            continue
        if not value or not str(value).strip():
            skipped += 1
            continue

        flat[key] = str(value).strip()
        translated += 1

    nested = build_nested(flat)

    out_path = rf'd:\killersudoku\game-src\utils\locales\{lang}.json'
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(nested, f, ensure_ascii=False, indent=2)

    print(f'✅ 已生成 {lang}.json')
    print(f'   翻译词条: {translated}')
    print(f'   未翻译（跳过）: {skipped}')
    print(f'   输出路径: {out_path}')

if __name__ == '__main__':
    main()
